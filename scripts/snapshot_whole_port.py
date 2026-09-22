#!/usr/bin/env python3
"""Read Whole Port!B2:C9 once per Shanghai day; append immutable snapshots.

Amounts and source URLs must never appear in logs or version control.
This private history database is independent of market.db and company.db.
"""
from __future__ import annotations

import argparse
import datetime as dt
import io
import json
import math
from pathlib import Path
import sqlite3
import sys
from zoneinfo import ZoneInfo

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from portfolio.holdings.sheet_book import SheetBookError, download_sheet_xlsx

TIMEZONE = ZoneInfo("Asia/Shanghai")
TAB = "Whole Port"
CELL_RANGE = "B2:C9"


def parse_whole_port(payload: bytes) -> list:
    import openpyxl

    try:
        wb = openpyxl.load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
        try:
            if TAB not in wb.sheetnames:
                raise SheetBookError("Whole Port tab missing")
            rows = [list(row) for row in wb[TAB].iter_rows(
                min_row=2, max_row=9, min_col=2, max_col=3, values_only=True)]
        finally:
            wb.close()
    except SheetBookError:
        raise
    except Exception as exc:
        raise SheetBookError("invalid workbook: " + type(exc).__name__) from None

    if len(rows) != 8 or rows[0][0] is not None or not isinstance(rows[0][1], str):
        raise SheetBookError("Whole Port B2:C2 header changed")
    if not rows[0][1].strip() or rows[0][1].startswith("#"):
        raise SheetBookError("Whole Port header missing or invalid")
    labels = []
    for i, (label, value) in enumerate(rows[1:], 3):
        if not isinstance(label, str) or not label.strip() or label.startswith("#"):
            raise SheetBookError("invalid label at B%d" % i)
        labels.append(label.strip())
        if isinstance(value, bool) or not isinstance(value, (float, int)) or not math.isfinite(value):
            raise SheetBookError("missing/non-numeric/error value at C%d" % i)
    if len(set(labels)) != 7 or labels[-2:] != ["total", "Net Total"]:
        raise SheetBookError("Whole Port total labels changed or duplicate labels")
    return rows


def capture(db_path: Path, fetch=download_sheet_xlsx,
            clock=lambda: dt.datetime.now(TIMEZONE)) -> str:
    db_path = Path(db_path)
    db_path.parent.mkdir(parents=True, exist_ok=True)
    with sqlite3.connect(db_path, timeout=60) as conn:
        conn.execute("""CREATE TABLE IF NOT EXISTS whole_port_daily (
            snapshot_date TEXT PRIMARY KEY,
            captured_at TEXT NOT NULL,
            timezone TEXT NOT NULL,
            sheet_tab TEXT NOT NULL,
            cell_range TEXT NOT NULL,
            cells_json TEXT NOT NULL
        )""")
        today = clock().astimezone(TIMEZONE).date().isoformat()
        if conn.execute("SELECT 1 FROM whole_port_daily WHERE snapshot_date=?", (today,)).fetchone():
            return "already_recorded"
        rows = parse_whole_port(fetch())
        captured = clock().astimezone(TIMEZONE)
        # Date is the actual observation day, including a fetch crossing midnight.
        cursor = conn.execute("""INSERT INTO whole_port_daily
            VALUES (?, ?, ?, ?, ?, ?) ON CONFLICT(snapshot_date) DO NOTHING""", (
                captured.date().isoformat(), captured.isoformat(timespec="seconds"),
                "Asia/Shanghai", TAB, CELL_RANGE,
                json.dumps(rows, ensure_ascii=False, allow_nan=False),
            ))
        return "recorded" if cursor.rowcount == 1 else "already_recorded"


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--db", type=Path, required=True,
                        help="Private cloud history database (never market.db/company.db)")
    args = parser.parse_args()
    if args.db.name in {"market.db", "company.db"}:
        parser.error("use a separate history database")
    try:
        status = capture(args.db)
    except Exception as exc:
        # Never print arbitrary exception text: upstream can include the secret URL.
        print("whole_port_snapshot failed: " + type(exc).__name__, file=sys.stderr)
        return 1
    print("whole_port_snapshot " + status)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
