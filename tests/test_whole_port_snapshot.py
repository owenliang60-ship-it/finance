import datetime as dt
import io
import json
import sqlite3
from concurrent.futures import ThreadPoolExecutor
from threading import Barrier

import openpyxl
import pytest

from scripts import snapshot_whole_port as job
from portfolio.holdings.sheet_book import SheetBookError


def workbook(bad=None):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Whole Port"
    for i, row in enumerate([
        [None, "USD"], ["美股", 100], ["BTC", 20], ["独立团资金", 0],
        ["A股", 30], ["现金", -10], ["total", 140], ["Net Total", 120],
    ], 2):
        ws.cell(i, 2, row[0])
        ws.cell(i, 3, row[1])
    if bad:
        ws[bad[0]] = bad[1]
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def test_exact_range_and_cached_values():
    rows = job.parse_whole_port(workbook())
    assert rows == [[None, "USD"], ["美股", 100], ["BTC", 20],
                    ["独立团资金", 0], ["A股", 30], ["现金", -10],
                    ["total", 140], ["Net Total", 120]]


@pytest.mark.parametrize("bad", [("C3", "#N/A"), ("C9", None),
                                  ("B9", "wrong"), ("C4", True),
                                  ("C5", "=1+2"), ("B6", "BTC")])
def test_bad_values_fail_closed(bad):
    with pytest.raises(SheetBookError):
        job.parse_whole_port(workbook(bad))


def test_bad_download():
    with pytest.raises(SheetBookError):
        job.parse_whole_port(b"<html>sign in</html>")


def test_daily_first_success_wins_and_uses_shanghai_date(tmp_path):
    db = tmp_path / "history.db"
    t = dt.datetime(2026, 9, 22, 17, tzinfo=dt.timezone.utc)
    assert job.capture(db, fetch=lambda: workbook(), clock=lambda: t) == "recorded"
    def must_not_fetch():
        raise AssertionError("duplicate day must not fetch")
    assert job.capture(db, fetch=must_not_fetch, clock=lambda: t) == "already_recorded"
    with sqlite3.connect(db) as conn:
        row = conn.execute("SELECT snapshot_date, captured_at, cells_json FROM whole_port_daily").fetchone()
    assert row[0] == "2026-09-23"
    assert row[1] == "2026-09-23T01:00:00+08:00"
    assert json.loads(row[2])[-1] == ["Net Total", 120]
    assert job.capture(db, fetch=lambda: workbook(), clock=lambda: t + dt.timedelta(days=1)) == "recorded"
    with sqlite3.connect(db) as conn:
        assert conn.execute("SELECT COUNT(*) FROM whole_port_daily").fetchone()[0] == 2


def test_failed_fetch_does_not_claim_day(tmp_path):
    db = tmp_path / "history.db"
    with pytest.raises(SheetBookError):
        job.capture(db, fetch=lambda: workbook(("C9", "#REF!")))
    with sqlite3.connect(db) as conn:
        assert conn.execute("SELECT COUNT(*) FROM whole_port_daily").fetchone()[0] == 0
    assert job.capture(db, fetch=lambda: workbook()) == "recorded"


def test_concurrent_runs_record_one_complete_snapshot(tmp_path):
    db = tmp_path / "history.db"
    barrier = Barrier(2)
    payload = workbook()
    def fetch():
        barrier.wait(timeout=10)
        return payload
    with ThreadPoolExecutor(max_workers=2) as pool:
        futures = [pool.submit(job.capture, db, fetch) for _ in range(2)]
        assert sorted(f.result() for f in futures) == ["already_recorded", "recorded"]
    with sqlite3.connect(db) as conn:
        rows = conn.execute("SELECT cells_json FROM whole_port_daily").fetchall()
        assert len(rows) == 1
        assert len(json.loads(rows[0][0])) == 8


def test_missing_tab():
    wb = openpyxl.Workbook()
    out = io.BytesIO()
    wb.save(out)
    with pytest.raises(SheetBookError, match="tab missing"):
        job.parse_whole_port(out.getvalue())


def test_cli_redacts_error_details(monkeypatch, tmp_path, capsys):
    def fail(*args, **kwargs):
        raise RuntimeError("private-sheet-url-and-balance")
    monkeypatch.setattr(job, "capture", fail)
    monkeypatch.setattr("sys.argv", ["snapshot", "--db", str(tmp_path / "history.db")])
    assert job.main() == 1
    assert capsys.readouterr().err == "whole_port_snapshot failed: RuntimeError\n"
